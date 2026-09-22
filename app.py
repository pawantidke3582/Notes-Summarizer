import os
import io
from flask import Flask, request, jsonify, render_template, send_from_directory
from flask_cors import CORS
from dotenv import load_dotenv
import pypdf
import docx
from groq import Groq

# Load environment variables
load_dotenv()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, 'templates') if os.path.exists(os.path.join(BASE_DIR, 'templates')) else os.path.join(os.getcwd(), 'templates')

STATIC_DIR = os.path.join(BASE_DIR, 'static')
if not os.path.exists(STATIC_DIR):
    STATIC_DIR = os.path.join(BASE_DIR, 'public', 'static')
if not os.path.exists(STATIC_DIR):
    STATIC_DIR = os.path.join(os.getcwd(), 'static')
if not os.path.exists(STATIC_DIR):
    STATIC_DIR = os.path.join(os.getcwd(), 'public', 'static')

app = Flask(
    __name__,
    template_folder=TEMPLATES_DIR,
    static_folder=STATIC_DIR
)
CORS(app)

# Configuration
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024  # 20 MB max file size
DEFAULT_MODEL = os.getenv("MODEL_NAME", "qwen/qwen3.8-27b")
FALLBACK_MODEL = "groq/compound-mini"


def get_groq_client(custom_api_key=None):
    """Initializes and returns Groq client using custom key or .env key."""
    load_dotenv(override=True)
    api_key = (custom_api_key or "").strip() or os.getenv("GROQ_API_KEY", "").strip()
    if not api_key or api_key == "your_groq_api_key_here":
        return None, "Groq API key not found. Please provide an API key in settings or set GROQ_API_KEY in your .env file."
    try:
        client = Groq(api_key=api_key)
        return client, None
    except Exception as e:
        return None, f"Failed to initialize Groq client: {str(e)}"


def extract_text_from_file(file_storage):
    """Extracts plain text from uploaded PDF, DOCX, or TXT file."""
    filename = file_storage.filename or ""
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    file_bytes = file_storage.read()
    text = ""

    if ext == 'pdf':
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            extracted_pages = []
            for page_num, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    extracted_pages.append(page_text)
            text = "\n\n".join(extracted_pages)
            if not text.strip():
                return None, "No readable text found in this PDF. It might be scanned or image-only."
        except Exception as e:
            return None, f"Error parsing PDF file: {str(e)}"

    elif ext in ['docx', 'doc']:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n\n".join(paragraphs)
            if not text.strip():
                return None, "No readable text found in this DOCX document."
        except Exception as e:
            return None, f"Error parsing DOCX file: {str(e)}"

    elif ext == 'txt':
        try:
            # Try utf-8 first, fallback to latin-1
            text = file_bytes.decode('utf-8')
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode('latin-1')
            except Exception as e:
                return None, f"Error decoding text file: {str(e)}"
    else:
        return None, f"Unsupported file type: .{ext}. Please upload a PDF, DOCX, or TXT file."

    cleaned_text = text.strip()
    if not cleaned_text:
        return None, "The uploaded file contains empty content."

    return cleaned_text, None


@app.route('/')
@app.route('/index.html')
@app.route('/api')
@app.route('/api/index')
@app.route('/api/index.py')
def index():
    """Serves the main application page."""
    if request.args.get('debug') == '1':
        return jsonify({
            "path": request.path,
            "url": request.url,
            "environ_PATH_INFO": request.environ.get('PATH_INFO'),
            "environ_SCRIPT_NAME": request.environ.get('SCRIPT_NAME'),
            "headers": dict(request.headers)
        })
    return render_template('index.html')


@app.route('/static/<path:filename>')
def serve_static(filename):
    """Serves static assets explicitly with multiple directory fallbacks."""
    for directory in [
        STATIC_DIR,
        os.path.join(BASE_DIR, 'public', 'static'),
        os.path.join(BASE_DIR, 'public'),
        os.path.join(os.getcwd(), 'static'),
        os.path.join(os.getcwd(), 'public', 'static')
    ]:
        if os.path.exists(os.path.join(directory, filename)):
            return send_from_directory(directory, filename)
    return send_from_directory(STATIC_DIR, filename)


@app.route('/css/<path:filename>')
def serve_css(filename):
    for directory in [
        os.path.join(BASE_DIR, 'public', 'css'),
        os.path.join(STATIC_DIR, 'css'),
        os.path.join(os.getcwd(), 'public', 'css'),
        os.path.join(os.getcwd(), 'static', 'css')
    ]:
        if os.path.exists(os.path.join(directory, filename)):
            return send_from_directory(directory, filename)
    return send_from_directory(os.path.join(STATIC_DIR, 'css'), filename)


@app.route('/js/<path:filename>')
def serve_js(filename):
    for directory in [
        os.path.join(BASE_DIR, 'public', 'js'),
        os.path.join(STATIC_DIR, 'js'),
        os.path.join(os.getcwd(), 'public', 'js'),
        os.path.join(os.getcwd(), 'static', 'js')
    ]:
        if os.path.exists(os.path.join(directory, filename)):
            return send_from_directory(directory, filename)
    return send_from_directory(os.path.join(STATIC_DIR, 'js'), filename)


@app.route('/api/status', methods=['GET'])
def status():
    """Checks service and configuration status."""
    load_dotenv(override=True)
    env_key = os.getenv("GROQ_API_KEY", "").strip()
    has_key = bool(env_key and env_key != "your_groq_api_key_here")
    return jsonify({
        "status": "online",
        "has_env_key": has_key,
        "default_model": DEFAULT_MODEL
    })


@app.route('/api/upload', methods=['POST'])
def upload_file():
    """Accepts a document file and returns extracted text and stats."""
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    file = request.files['file']
    if not file or file.filename == '':
        return jsonify({"error": "No file selected"}), 400

    text, error = extract_text_from_file(file)
    if error:
        return jsonify({"error": error}), 400

    words = len(text.split())
    chars = len(text)

    return jsonify({
        "filename": file.filename,
        "text": text,
        "word_count": words,
        "char_count": chars,
        "message": f"Successfully extracted {words:,} words from {file.filename}."
    })


@app.route('/api/summarize', methods=['POST'])
def summarize():
    """Generates AI summary using specified mode and Groq API."""
    data = request.get_json() or {}
    text = data.get("text", "").strip()
    mode = data.get("mode", "short").lower()
    custom_api_key = data.get("api_key", "").strip()

    if not text:
        return jsonify({"error": "Please provide notes text or upload a document to summarize."}), 400

    client, err = get_groq_client(custom_api_key)
    if err:
        return jsonify({"error": err}), 401

    # Prompt configurations
    prompts = {
        "short": (
            "You are an expert notes summarizer. Provide a concise, clear, and easy-to-understand "
            "executive summary of the provided notes in 2 to 3 sentences. Follow it with exactly 3 "
            "high-impact bullet points capturing the key highlights. Keep the language natural and engaging."
        ),
        "detailed": (
            "You are an expert academic and professional notes summarizer. Provide an organized, in-depth, "
            "and comprehensive summary of the provided notes. Structure your response with:\n"
            "### 📌 Executive Overview\n"
            "### 🔍 Detailed Breakdown & Core Concepts\n"
            "### 💡 Key Insights & Examples\n"
            "### 🏁 Conclusion & Summary Takeaway\n"
            "Use clear Markdown formatting, bolded key terms, and bullet points."
        ),
        "points": (
            "You are an expert study assistant. Extract all the most critical and important points, "
            "facts, principles, definitions, or instructions from the notes. Present them as a well-categorized, "
            "easy-to-read bulleted list. Bold the primary keywords in each point for quick reading."
        ),
        "questions": (
            "You are an educational tutor. Based STRICTLY on the provided notes, create 5 to 7 essential "
            "review/study questions that test understanding of the material. For each question, provide a clear, "
            "accurate answer directly derived from the notes. Format each item as:\n"
            "**Q[number]: [Question]?**\n"
            "> **Answer:** [Clear concise answer grounded in notes]\n"
        )
    }

    system_instruction = prompts.get(mode, prompts["short"])

    user_content = f"Here are the notes to summarize:\n\n---\n{text}\n---"

    try:
        completion = client.chat.completions.create(
            model=DEFAULT_MODEL,
            messages=[
                {"role": "system", "content": system_instruction},
                {"role": "user", "content": user_content}
            ],
            temperature=0.3,
            max_tokens=2048,
        )
        summary = completion.choices[0].message.content
        return jsonify({
            "summary": summary,
            "mode": mode,
            "model": DEFAULT_MODEL
        })
    except Exception as e:
        # Fallback to secondary fast model if primary fails
        try:
            completion = client.chat.completions.create(
                model=FALLBACK_MODEL,
                messages=[
                    {"role": "system", "content": system_instruction},
                    {"role": "user", "content": user_content}
                ],
                temperature=0.3,
                max_tokens=2048,
            )
            return jsonify({
                "summary": completion.choices[0].message.content,
                "mode": mode,
                "model": FALLBACK_MODEL
            })
        except Exception as fallback_err:
            return jsonify({"error": f"Groq API Error: {str(e)}"}), 500


@app.route('/api/chat', methods=['POST'])
def chat():
    """Answers user questions strictly grounded in the uploaded notes."""
    data = request.get_json() or {}
    notes_text = data.get("notes", "").strip()
    question = data.get("question", "").strip()
    history = data.get("history", [])
    custom_api_key = data.get("api_key", "").strip()

    if not notes_text:
        return jsonify({"error": "No notes found. Please upload a document or paste notes first."}), 400

    if not question:
        return jsonify({"error": "Please enter a question to ask."}), 400

    client, err = get_groq_client(custom_api_key)
    if err:
        return jsonify({"error": err}), 401

    system_prompt = (
        "You are an intelligent, precise AI study assistant. Your task is to answer the user's questions "
        "ONLY using the provided notes context below.\n\n"
        "STRICT GROUNDING RULES:\n"
        "1. Answer strictly based on the provided notes. Do NOT use outside knowledge or hallucinate facts.\n"
        "2. If the answer cannot be found or directly deduced from the provided notes, you MUST state exactly:\n"
        "   \"I cannot find this information in your notes. Please make sure your notes cover this topic or rephrase your question.\"\n"
        "3. Keep answers clear, accurate, concise, and formatted with Markdown (bullet points, bold text, code blocks) where helpful.\n\n"
        f"--- NOTES CONTEXT ---\n{notes_text}\n--- END NOTES CONTEXT ---"
    )

    messages = [{"role": "system", "content": system_prompt}]

    # Append recent conversation history (limit to last 6 turns for context window)
    for msg in history[-6:]:
        role = msg.get("role")
        content = msg.get("content")
        if role in ["user", "assistant"] and content:
            messages.append({"role": role, "content": content})

    messages.append({"role": "user", "content": question})

    try:
        completion = client.chat.completions.create(
            model=DEFAULT_MODEL,
            messages=messages,
            temperature=0.2,
            max_tokens=1024,
        )
        reply = completion.choices[0].message.content
        return jsonify({"reply": reply})
    except Exception as e:
        try:
            completion = client.chat.completions.create(
                model=FALLBACK_MODEL,
                messages=messages,
                temperature=0.2,
                max_tokens=1024,
            )
            return jsonify({"reply": completion.choices[0].message.content})
        except Exception as fallback_err:
            return jsonify({"error": f"Groq API Error: {str(e)}"}), 500


if __name__ == '__main__':
    port = int(os.getenv("PORT", 5000))
    debug = os.getenv("DEBUG", "True").lower() == "true"
    print(f" * AI Notes Summarizer Chatbot starting on http://127.0.0.1:{port}")
    app.run(host='0.0.0.0', port=port, debug=debug)
